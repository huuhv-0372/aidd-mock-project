"""Slide 16 — 提案の要諦 as four point cards, and the gate that keeps it honest.

The chapter has two shapes in the wild: a Japanese proposal writes it as a
comparison matrix (slide 21), an English one as `###` subsections (slide 16).
These tests pin that both reach a bespoke slide, that the fourth point is not
lost to a three-card design, and — the part that matters most — that every text
slot on slide 16 is bound, because an unbound slot ships another client's
Japanese sample sentence to this client's deck.
"""
from __future__ import annotations

import subprocess
import sys

import pytest

pytest.importorskip('pptx')

from lib.profile_parser import parse_profile  # noqa: E402
from lib.templates.svn import CHAPTER_DIVIDERS, SECTION_TO_SLIDES, SVNProposalTemplate  # noqa: E402

POINTS_DOC = '''# Proposal

## 3. Essence of Sun*'s Proposal

### 3.1. Proven Global Team

- Existing engagement since June 2025, with continuous context carried forward.
- Over 500 clients delivered globally with consistent QCD outcomes.

### 3.2. AI-Driven Development

Sun* applies Claude Code across the lifecycle. It is a default workflow, not assistance.

### 3.3. Hardware Integration

Sun* has delivered BLE and OTA projects. The architecture reflects those lessons.

### 3.4. Wellness Domain Context

- Understanding of GDPR-equivalent obligations without over-engineering.
- Awareness of the wellness versus medical device distinction.
'''

MATRIX_DOC = '''# Proposal

## 提案の要諦

| 観点 | A案 | B案 |
|---|---|---|
| コスト | 低 | 高 |
| 期間 | 長 | 短 |
'''


def test_subsections_become_four_point_cards():
    points = parse_profile(POINTS_DOC).approach_points
    assert len(points) == 4
    assert [p['title'] for p in points] == [
        'Proven Global Team', 'AI-Driven Development',
        'Hardware Integration', 'Wellness Domain Context',
    ]


def test_the_headline_slot_is_the_points_own_first_sentence():
    """Nothing is invented for the card's third slot."""
    card = parse_profile(POINTS_DOC).approach_points[1]
    assert card['subtitle'] == 'Sun* applies Claude Code across the lifecycle.'
    assert card['body'] == 'It is a default workflow, not assistance.'


def test_headline_and_body_never_print_the_same_sentence_twice():
    for card in parse_profile(POINTS_DOC).approach_points:
        assert card['subtitle'] not in card['body']


def test_list_only_subsections_still_produce_a_card():
    """3.1 and 3.4 are nothing but bullets; `parse_prose` alone would drop them."""
    cards = {p['title']: p for p in parse_profile(POINTS_DOC).approach_points}
    assert 'Existing engagement' in cards['Proven Global Team']['subtitle']
    assert '500 clients' in cards['Proven Global Team']['body']


def test_a_matrix_chapter_still_fills_the_table_slide_and_no_cards():
    profile = parse_profile(MATRIX_DOC)
    assert profile.approach_comparison, 'the matrix must still reach slide 21'
    assert profile.approach_points == [], 'a table is not a set of point cards'


def test_both_slides_belong_to_the_chapter_and_its_divider():
    assert SECTION_TO_SLIDES['approach_comparison'] == [16, 21]
    assert CHAPTER_DIVIDERS[14] == [16, 21]


def test_a_fifth_point_is_reported_not_silently_dropped(capsys):
    doc = POINTS_DOC + '\n### 3.5. Fifth Point\n\nA fifth point that has no card.\n'
    points = parse_profile(doc).approach_points
    assert len(points) == 4
    assert 'stay in the markdown only' in capsys.readouterr().out


def test_every_text_slot_on_slide_16_is_bound_by_a_role():
    """Twelve card slots, four cards, no slot left holding template sample text."""
    roles = {r.name: r for r in SVNProposalTemplate().slide_role_configs[16].roles}
    pairs = roles['essence_points'].find_params['pair_names']
    headlines = roles['essence_headlines'].find_params['names']
    assert len(pairs) == 4 and len(headlines) == 4
    assert roles['essence_headlines'].item_field == 'subtitle'
    bound = {n for pair in pairs for n in pair} | set(headlines)
    assert len(bound) == 12, 'title + headline + body for each of four cards'


def test_the_template_contract_gate_passes(scripts_dir):
    """And it is a real gate — removing a role makes it fail (verified manually)."""
    result = subprocess.run(
        [sys.executable, str(scripts_dir / 'verify-template-contract.py')],
        capture_output=True, text=True, cwd=str(scripts_dir),
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert 'SAMPLE_TEXT_SLIDE_16' in result.stdout
    assert 'no unbound sample text' in result.stdout


def test_code_fences_and_rules_never_reach_a_card():
    """An ASCII diagram is layout; half of one on a slide reads as damage."""
    doc = '''# P

## 提案の要諦

### 1. With a diagram

Intro sentence here.

```
step one
   -> step two
```

Closing sentence here.

---
'''
    card = parse_profile(doc).approach_points[0]
    combined = card['subtitle'] + card['body']
    assert 'step one' not in combined
    assert '```' not in combined
    assert not combined.rstrip().endswith('---')


# --- One size, one baseline: the card row as rendered ------------------------

CARD_NAMES = {
    'titles': ['Google Shape;1002;p16', 'Google Shape;1008;p16',
               'Google Shape;1014;p16', 'Google Shape;1020;p16'],
    'headlines': ['Google Shape;1003;p16', 'Google Shape;1009;p16',
                  'Google Shape;1015;p16', 'Google Shape;1021;p16'],
    'bodies': ['Google Shape;1004;p16', 'Google Shape;1010;p16',
               'Google Shape;1016;p16', 'Google Shape;1022;p16'],
}

# Prose the length a real proposal writes it — the four cards' text differs in
# length by a factor of two, which is what made the renderer size each box on
# its own.
UNEVEN_DOC = '''# Proposal

## 3. Essence of Sun*'s Proposal

### 3.1. Proven Global Team

- Existing engagement with the client since June 2025, with continuous context carried forward.
- Over 500 clients delivered globally with consistent QCD outcomes.

### 3.2. AI-Driven Development as a First-Class Methodology

Sun* applies Claude Code and the Claude API across the full development lifecycle: This is not "AI assistance" — it is a default workflow with human review gates at every critical transition.

### 3.3. Hardware Integration

Sun* has delivered projects involving BLE peripherals, firmware OTA update flows, and high-throughput time-series ingestion. The architecture reflects lessons learned in those engagements.

### 3.4. Wellness Domain Context

- Understanding of GDPR-equivalent privacy obligations without over-engineering for HIPAA or medical device rules that do not apply to wellness products.
- Awareness of the regulatory distinction between "wellness" and "medical device" claims.
'''


@pytest.fixture(scope='module')
def card_slide(tmp_path_factory, template_path):
    """Slide 16 of a rendered deck, as {shape name: shape}."""
    from pptx import Presentation
    from lib.renderer import PPTXRenderer

    profile = parse_profile(UNEVEN_DOC)
    out_dir = tmp_path_factory.mktemp('cards')
    out = PPTXRenderer().render_from_profile(
        profile, str(template_path), 'cards', str(out_dir),
        extra_slides=list(profile.auto_extra_slides),
    )

    def walk(shapes):
        for shape in shapes:
            yield shape
            if 'GROUP' in str(shape.shape_type):
                yield from walk(shape.shapes)

    for slide in Presentation(str(out)).slides:
        shapes = {s.name: s for s in walk(slide.shapes)}
        if 'Google Shape;1003;p16' in shapes:
            return shapes
    pytest.fail('the four-card slide is not in the deck')


def _sizes(shapes, names):
    return {shapes[name].text_frame.paragraphs[0].runs[0].font.size for name in names}


def test_headlines_and_bodies_are_set_at_one_size(card_slide):
    """The bold headline and the plain body under it are authored the same size."""
    sizes = _sizes(card_slide, CARD_NAMES['headlines'] + CARD_NAMES['bodies'])
    assert len(sizes) == 1, 'one size for the whole card row, not one per box'
    assert sizes.pop() is not None, 'the size is written, never left to the renderer'


def test_no_card_box_is_left_on_auto_fit(card_slide):
    """`normAutofit` is the renderer sizing each box on its own — the defect."""
    from pptx.oxml.ns import qn

    for name in CARD_NAMES['headlines'] + CARD_NAMES['bodies']:
        bodyPr = card_slide[name].text_frame._txBody.find(qn('a:bodyPr'))
        assert bodyPr.find(qn('a:noAutofit')) is not None, name
        assert bodyPr.find(qn('a:normAutofit')) is None, name
        assert bodyPr.find(qn('a:spAutoFit')) is None, name


def test_the_row_shares_one_top_and_one_height(card_slide):
    """Four cards, one baseline — for the headline row and the body row alike."""
    for row in ('headlines', 'bodies'):
        boxes = [card_slide[name] for name in CARD_NAMES[row]]
        assert len({box.top for box in boxes}) == 1, row
        assert len({box.height for box in boxes}) == 1, row


def test_card_text_stays_inside_the_card_it_is_drawn_in(card_slide):
    """Re-dividing the band may never push text past the card's own panel."""
    panels = ['Google Shape;1001;p16', 'Google Shape;1007;p16',
              'Google Shape;1013;p16', 'Google Shape;1019;p16']
    for panel_name, body_name in zip(panels, CARD_NAMES['bodies']):
        panel, body = card_slide[panel_name], card_slide[body_name]
        assert body.top + body.height <= panel.top + panel.height, body_name


def test_the_title_is_never_smaller_than_the_text_it_heads(card_slide):
    title = min(card_slide[n].text_frame.paragraphs[0].runs[0].font.size
                for n in CARD_NAMES['titles'])
    body = _sizes(card_slide, CARD_NAMES['bodies']).pop()
    assert title >= body


def test_a_body_slot_opens_a_sentence_rather_than_continuing_one(card_slide):
    """The card whose prose carries a mid-sentence dash — see split_headline."""
    body = card_slide['Google Shape;1010;p16'].text_frame.text.strip()
    assert body.startswith('This is not'), body


def test_a_pruned_slides_overflow_does_not_become_an_appendix_slide(tmp_path,
                                                                    template_path):
    """Slide 15 is filled and then dropped; its leftovers must go with it.

    Every configured slide is filled, including the ones no section claims —
    slide 15 draws the same 提案の要諦 cards as slide 16, in three narrower
    slots. Both truncate into the same queue, and slide 15's fragments used to
    outlive the slide itself: the deck ended up with an appendix slide holding
    `s" and "medical device" claims.`, a piece cut from a sentence printed in
    full two slides earlier, by a slide the reader never saw.
    """
    from lib.profile_parser import parse_profile
    from lib.renderer import PPTXRenderer

    doc = (
        '# P\n\n## 1. System Overview\n\n### 1.1. Project Background\n\nBackground.\n\n'
        "## 3. Essence of Sun*'s Proposal\n\n"
        + ''.join(
            f'### 3.{i}. Point number {i}\n\n'
            f'Headline sentence for point {i}. ' + ('Elaboration clause. ' * 18) + '\n\n'
            for i in range(1, 5))
    )
    profile = parse_profile(doc)
    renderer = PPTXRenderer()
    out = renderer.render_from_profile(profile, str(template_path), 'pruned-overflow',
                                       str(tmp_path),
                                       extra_slides=list(profile.auto_extra_slides))

    queued = renderer._overflow_extras.get(('approach_comparison', 'essence_points'), [])
    assert {slide for slide, _ in queued} >= {15, 16}, (
        'the test no longer exercises both the pruned twin and the kept slide')

    from pptx import Presentation
    prs = Presentation(str(out))
    printed = {
        ' '.join(shape.text_frame.text.split())
        for slide in prs.slides for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip()
    }
    from_pruned = [text for slide, text in queued if slide == 15]
    assert from_pruned, 'slide 15 truncated nothing — widen the fixture'
    for text in from_pruned:
        if text in [t for s, t in queued if s == 16]:
            continue  # slide 16 cut the same place; that one is legitimate
        assert ' '.join(text.split()) not in printed, (
            f'a fragment from the pruned slide 15 reached the deck: {text[:60]!r}')


def _all_text(pptx_path) -> list:
    """Every string the deck shows, table cells included.

    A table's cell text lives on the cell, not on the shape's own text frame,
    so a check that reads only text frames cannot see a matrix at all.
    """
    from pptx import Presentation
    out = []
    for slide in Presentation(str(pptx_path)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(' '.join(shape.text_frame.text.split()))
            if getattr(shape, 'has_table', False) and shape.has_table:
                for row in shape.table.rows:
                    out.extend(' '.join(c.text.split()) for c in row.cells)
    return out


def test_a_matrix_ONLY_chapter_does_not_also_ship_the_empty_card_slide(
        template_path, tmp_path):
    """Slides 16 and 21 are ALTERNATIVE designs for one chapter, not two pages.

    16 draws numbered point cards from `approach_points`; 21 draws the matrix
    from `approach_comparison`. A chapter written as a table only fills 21, and
    16 used to survive anyway — the prune pass kept the first slide of every
    section unconditionally — reaching the client as a page carrying nothing
    but its own header and four bare badges.
    """
    from pptx import Presentation

    from lib.renderer import PPTXRenderer

    doc = """# システム開発ご提案書

## 1. システムの概要 — System Overview

### 1.1 プロジェクト背景

現行業務の課題を解決するための提案です。

## 提案の要諦 — Essence of the Proposal

| 比較項目 | 弊社案 | 他社一般 |
|---|---|---|
| 開発体制 | 専任チーム | 都度アサイン |
| 品質保証 | 自動テスト | 手動のみ |
"""
    profile = parse_profile(doc)
    assert profile.approach_points == [], 'the fixture must fill the matrix only'
    assert profile.approach_comparison, 'the matrix itself must be parsed'

    out = PPTXRenderer().render_from_profile(
        profile, str(template_path), 'matrix-only', str(tmp_path),
        extra_slides=list(profile.auto_extra_slides))

    texts = _all_text(out)
    assert not any(text == '本提案のポイント' for text in texts), (
        "slide 16's header reached the deck, so the empty card slide shipped")
    assert any('比較項目' in text or '専任チーム' in text for text in texts), \
        'the matrix the chapter DID write must still be there'
