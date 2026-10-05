"""Appendix chapters borrowing a real template design instead of a drawn one.

`extra_slide_layouts.py` redraws, from primitives, designs this deck already
ships — its `numbered_points` is a worse slide 16. The donor engine lends the
real slide instead, and these tests pin the three things that make that safe:
the donor must survive pruning, it must fit every item, and a chapter that
matches nothing must still render exactly as it did before.
"""
from __future__ import annotations

import pytest

pytest.importorskip('pptx')

from pptx import Presentation  # noqa: E402

from lib.extra_slide_classifier import build_extra_slide_entry  # noqa: E402
from lib.profile_parser import parse_profile  # noqa: E402
from lib.renderer import PPTXRenderer  # noqa: E402
from lib.templates.svn import DONOR_SLIDES  # noqa: E402

FIXTURE = 'appendix-chapter-borrows-a-design.md'


def _subsections(n: int) -> dict:
    return {f'{i}. Point {i}': f'Headline {i}. Elaboration for point {i}.'
            for i in range(1, n + 1)}


def _render(tmp_path, template_path, markdown: str, name: str = 'donor-test'):
    """Render the way `gen-slide.py` does — auto extras included.

    The auto-routed chapters live on the profile and are handed to the renderer
    as `extra_slides`; rendering without them exercises a path no entry point
    uses.
    """
    profile = parse_profile(markdown)
    renderer = PPTXRenderer()
    out = renderer.render_from_profile(profile, str(template_path), name, str(tmp_path),
                                       extra_slides=list(profile.auto_extra_slides))
    return Presentation(str(out))


def _walk(shapes):
    for shape in shapes:
        yield shape
        if 'GROUP' in str(shape.shape_type):
            yield from _walk(shape.shapes)


def _texts(slide) -> list:
    return [' '.join(s.text_frame.text.split()) for s in _walk(slide.shapes)
            if s.has_text_frame and s.text_frame.text.strip()]


# --- classifier: when is a donor offered at all ----------------------------

def test_a_matching_chapter_is_offered_the_donor():
    entry = build_extra_slide_entry('Ways of Working', subsections=_subsections(4))
    assert entry['donor'] == 16
    assert len(entry['donor_items']) == 4
    assert entry['donor_items'][0]['subtitle'] == 'Headline 1.'
    assert entry['donor_items'][0]['body'] == 'Elaboration for point 1.'


def test_a_chapter_with_more_items_than_the_design_holds_is_NOT_offered_one():
    """Past the largest design in the library, no donor — and nothing dropped.

    The bound is the library's, not any one slide's: seven points exceed the
    six-row design (slide 38), which is the roomiest text layout the template
    ships. The chapter then falls to bullets, which never truncates.
    """
    entry = build_extra_slide_entry('Seven Ways', subsections=_subsections(7))
    assert 'donor' not in entry
    assert entry['layout'] == 'bullets'
    assert len(entry['bullets']) == 7, 'every point still reaches the deck'


def test_the_tightest_design_that_fits_is_the_one_offered():
    """A three-point chapter takes the three-card slide, not the six-row one."""
    assert build_extra_slide_entry('Three', subsections=_subsections(3))['donor'] == 15
    assert build_extra_slide_entry('Four', subsections=_subsections(4))['donor'] == 16
    assert build_extra_slide_entry('Five', subsections=_subsections(5))['donor'] == 38
    assert build_extra_slide_entry('Six', subsections=_subsections(6))['donor'] == 38


def test_a_six_point_chapter_reaches_the_deck_only_through_a_donor():
    """The drawn numbered_points layout caps at five; six is donor-only."""
    entry = build_extra_slide_entry('Six Ways', subsections=_subsections(6))
    assert entry['layout'] == 'numbered_points'
    assert len(entry['donor_items']) == 6


def test_a_table_chapter_is_not_offered_a_point_card_design():
    entry = build_extra_slide_entry(
        'Cost', table_rows=[{'Item': 'Dev', 'Amount': '1'}])
    assert 'donor' not in entry
    assert entry['layout'] == 'table'


def test_donor_titles_lose_the_documents_outline_numbering():
    entry = build_extra_slide_entry('Ways', subsections={
        '2.1. One Team': 'A sentence. Another sentence.',
        '2.2. Trunk Based': 'A sentence. Another sentence.',
        '2.3. Evidence First': 'A sentence. Another sentence.',
    })
    assert [p['title'] for p in entry['donor_items']] == [
        'One Team', 'Trunk Based', 'Evidence First']


def test_every_registered_donor_is_a_configured_slide_with_bound_roles():
    """A donor whose slots no role writes would lend its sample text along."""
    from lib.templates.svn import SVNProposalTemplate
    configs = SVNProposalTemplate().slide_role_configs
    for slide_no, spec in DONOR_SLIDES.items():
        assert slide_no in configs, f'donor {slide_no} has no role config'
        assert configs[slide_no].roles, f'donor {slide_no} binds no shapes'
        assert spec['capacity'] >= 1 and spec['source'] and spec['shape']


# --- end to end ------------------------------------------------------------

def test_a_borrowed_design_carries_the_chapters_own_content(tmp_path, template_path,
                                                            fixtures_dir):
    prs = _render(tmp_path, template_path, (fixtures_dir / FIXTURE).read_text('utf-8'))
    borrowed = [s for s in prs.slides if 'Appendix ｜ Ways of Working' in _texts(s)]
    assert len(borrowed) == 1, 'the chapter should occupy exactly one borrowed slide'
    texts = _texts(borrowed[0])
    for title in ('One Team, Two Timezones', 'Trunk-Based Delivery',
                  'Evidence Over Opinion', 'Handover From Day One'):
        assert title in texts
    assert 'Engineers sit in Da Nang and Osaka.' in texts
    # The design's own badges survive the transplant.
    for badge in ('01', '02', '03', '04'):
        assert badge in texts


def test_a_borrowed_design_ships_none_of_its_original_japanese(tmp_path, template_path,
                                                               fixtures_dir):
    """The donor was authored for another client. None of that may come along."""
    import re
    kana = re.compile(r'[ぁ-ゟ゠-ヿ]')
    prs = _render(tmp_path, template_path, (fixtures_dir / FIXTURE).read_text('utf-8'))
    borrowed = [s for s in prs.slides if 'Appendix ｜ Ways of Working' in _texts(s)][0]
    leftovers = [t for t in _texts(borrowed) if kana.search(t)]
    assert leftovers == [], leftovers


def test_the_donor_survives_the_prune_that_deletes_its_original(tmp_path, template_path,
                                                                fixtures_dir):
    """The whole hazard: a lendable slide is one pruning is about to delete.

    The fixture's chapters claim neither slide 16 nor its chapter divider, so
    the original is gone from the output — and the borrowed copy is still there.
    """
    prs = _render(tmp_path, template_path, (fixtures_dir / FIXTURE).read_text('utf-8'))
    all_texts = [t for s in prs.slides for t in _texts(s)]
    assert 'Appendix ｜ Ways of Working' in all_texts, 'the borrowed copy'
    assert '提案の要諦' not in all_texts, 'the original slide 16 should be pruned'


def test_a_chapter_matching_no_donor_renders_as_it_always_did(tmp_path, template_path,
                                                              fixtures_dir):
    """No donor must mean no change — the engine can only ever upgrade."""
    extra = ''.join(f'\n\nText.\n\n### 2.{n}. Way {n}' for n in range(5, 9))
    md = (fixtures_dir / FIXTURE).read_text('utf-8').replace(
        '### 2.4. Handover From Day One',
        '### 2.4. Handover From Day One' + extra)
    prs = _render(tmp_path, template_path, md, name='no-donor')
    chapter = [s for s in prs.slides if 'Appendix ｜ Ways of Working' in _texts(s)]
    assert len(chapter) == 1
    texts = _texts(chapter[0])
    # The donor's numbered badges are its signature; the drawn layout has none.
    assert '01' not in texts and '04' not in texts, 'this must be the drawn layout'
    # And the drawn layout shows every point — which is exactly why it won.
    joined = ' '.join(texts)
    for title in ('One Team', 'Trunk-Based', 'Evidence Over Opinion',
                  'Handover From Day One', 'Way 5', 'Way 8'):
        assert title in joined
