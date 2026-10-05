"""A drawn block is a picture, and the deck must treat it as one.

An architecture chapter carries its diagram as a fenced block of box-drawing
characters. Everything that makes such a block readable is geometry — one column
per character, one line per row — and the appendix path used to collapse it with
`' '.join(split())` and set the result in a proportional font with word wrap. The
diagram arrived as a paragraph of `│ │ ▼ ──▶` with its boxes gone.

These tests pin the three things that fixes it: the block is found, it is
recognised as a drawing rather than a listing, and it reaches the slide as an
image with its own lines intact.
"""
from __future__ import annotations

import pytest

pytest.importorskip('pptx')
pytest.importorskip('PIL')

from pptx import Presentation  # noqa: E402

from lib.diagram_image import _trimmed, find_mono_font, looks_like_diagram, render  # noqa: E402
from lib.extra_slide_classifier import build_extra_slide_entry  # noqa: E402
from lib.extra_slide_layouts import render_extra_layout  # noqa: E402
from lib.markdown_primitives import parse_fenced_blocks  # noqa: E402

SLIDE_W, SLIDE_H = 9144000, 5143500

PIPELINE = """Ring (firmware)
   │  BLE
   ▼
Mobile app (iOS / Android)
   │  HTTPS
   ▼
AWS backend"""

CHAPTER = f"""**Before (today)**: no integrated channel exists.

**After**:

```
{PIPELINE}
```
"""


def _blank():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    return prs.slides.add_slide(prs.slide_layouts[6])


# --- finding the block -----------------------------------------------------

def test_a_fenced_block_is_returned_with_its_lines_intact():
    (block,) = parse_fenced_blocks(CHAPTER)
    assert block.splitlines() == PIPELINE.splitlines()


def test_the_fence_markers_themselves_are_not_part_of_the_block():
    assert '```' not in parse_fenced_blocks(CHAPTER)[0]


def test_a_chapter_with_no_fence_yields_nothing():
    assert parse_fenced_blocks('Just prose.\n\nMore prose.') == []


def test_an_unclosed_fence_still_gives_up_its_diagram():
    """A truncated document should not cost the reader the whole drawing."""
    assert parse_fenced_blocks(f'text\n\n```\n{PIPELINE}') == [PIPELINE]


# --- drawing vs listing ----------------------------------------------------

def test_box_drawing_counts_as_a_diagram():
    assert looks_like_diagram(PIPELINE)


def test_a_yaml_listing_is_not_a_diagram():
    """It reads perfectly well as text; only a drawing loses meaning reflowed."""
    assert not looks_like_diagram('name: api\nport: 8080\nreplicas: 3')


def test_a_single_line_is_not_a_diagram():
    assert not looks_like_diagram('A ──▶ B')


# --- the block keeps its shape --------------------------------------------

def test_the_common_indent_goes_and_the_relative_one_stays():
    lines = _trimmed('    root\n        child\n    sibling')
    assert lines == ['root', '    child', 'sibling']


def test_blank_edges_are_dropped_but_inner_blank_lines_are_not():
    assert _trimmed('\n\nA\n\nB\n\n') == ['A', '', 'B']


# --- classification --------------------------------------------------------

def test_a_chapter_holding_a_drawing_becomes_a_diagram_slide():
    entry = build_extra_slide_entry('6.1. Architecture', raw_text=CHAPTER)
    assert entry['layout'] == 'diagram'
    assert entry['diagram'].splitlines() == PIPELINE.splitlines()


def test_the_chapters_prose_survives_as_the_lead_in():
    entry = build_extra_slide_entry('6.1. Architecture', raw_text=CHAPTER)
    assert any('no integrated channel exists' in line for line in entry['lead_in'])


def test_the_lead_in_carries_no_markdown_markers():
    """`**Before (today)**:` reached the deck with its asterisks showing."""
    entry = build_extra_slide_entry('6.1. Architecture', raw_text=CHAPTER)
    assert not any('**' in line for line in entry['lead_in'])


def test_a_chapter_whose_fence_is_a_listing_is_not_a_diagram_slide():
    chapter = 'Config we deploy:\n\n```\nname: api\nport: 8080\n```'
    assert build_extra_slide_entry('5.7. Config', raw_text=chapter)['layout'] != 'diagram'


# --- rendering -------------------------------------------------------------

@pytest.mark.skipif(not find_mono_font(), reason='no monospace face on this host')
def test_the_picture_fits_the_box_and_keeps_its_proportions():
    stream, w, h = render(PIPELINE, 8.0, 3.0)
    assert stream is not None
    assert 0 < w <= 8.0 and 0 < h <= 3.0
    # Wider than tall, as the block itself is.
    assert w > h


@pytest.mark.skipif(not find_mono_font(), reason='no monospace face on this host')
def test_the_diagram_reaches_the_slide_as_a_picture():
    slide = _blank()
    render_extra_layout(slide, build_extra_slide_entry('6.1. Architecture',
                                                       raw_text=CHAPTER),
                        SLIDE_W, SLIDE_H)
    pictures = [s for s in slide.shapes if s.shape_type == 13]
    assert len(pictures) == 1, 'the drawing is one image, not reflowed text'


@pytest.mark.skipif(not find_mono_font(), reason='no monospace face on this host')
def test_the_drawing_is_never_written_out_as_text_beside_its_picture():
    slide = _blank()
    render_extra_layout(slide, build_extra_slide_entry('6.1. Architecture',
                                                       raw_text=CHAPTER),
                        SLIDE_W, SLIDE_H)
    drawn = ' '.join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert '▼' not in drawn and 'Mobile app (iOS / Android)' not in drawn


def test_an_empty_diagram_does_not_raise():
    slide = _blank()
    render_extra_layout(slide, {'layout': 'diagram', 'title': 'X', 'diagram': '',
                                'lead_in': ['Only prose.']}, SLIDE_W, SLIDE_H)
    drawn = ' '.join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert 'Only prose.' in drawn


# --- a listing is not a drawing --------------------------------------------

SHELL_BLOCK = ("cat access.log | grep ERROR | awk '{print $1}' | sort | uniq -c\n"
               'tail -f app.log | grep WARN')
ASCII_BOX = ('+--------+      +--------+\n'
             '|  App   | ---> |  API   |\n'
             '+--------+      +--------+')


@pytest.mark.parametrize('block,what', [
    (SHELL_BLOCK, 'a shell pipeline'),
    ('const total = a + b + c;\nconst other = d + e || f;', 'javascript'),
    ("SELECT a || b FROM t\nWHERE x || y = 'z' || 'w'", 'sql'),
])
def test_punctuation_alone_does_not_make_a_drawing(block, what):
    """`+` and `|` are box art in a box and punctuation everywhere else.

    Counting them alone, four of either was enough: a shell pipeline read as a
    diagram, took the slide, and the chapter's table was dropped on the way —
    the one outcome the appendix routing exists to refuse.
    """
    assert not looks_like_diagram(block), f'{what} was read as a drawing'


def test_ascii_box_art_is_still_a_drawing():
    """The `+---` corner run is what separates art from arithmetic."""
    assert looks_like_diagram(ASCII_BOX)


def test_a_labelled_fence_is_a_listing_whatever_it_holds():
    """The author said `bash`. That settles it before any character is counted."""
    from lib.extra_slide_classifier import build_extra_slide_entry

    entry = build_extra_slide_entry('Runbook', raw_text=f'Intro.\n\n```bash\n{ASCII_BOX}\n```\n')
    assert entry['layout'] != 'diagram'


def test_an_unlabelled_fence_of_box_art_is_still_drawn():
    from lib.extra_slide_classifier import build_extra_slide_entry

    entry = build_extra_slide_entry('Architecture', raw_text=f'Intro.\n\n```\n{ASCII_BOX}\n```\n')
    assert entry['layout'] == 'diagram'


def test_a_diagram_chapter_keeps_its_table():
    """The drawing takes the slide; the table gets one of its own.

    A chapter is not reduced to one shape. Before, a table sharing a chapter
    with a diagram left the deck entirely and nothing said so.
    """
    from lib.profile_parser import parse_profile

    doc = (
        '# P\n\n## 1. System Overview\n\n### 1.1. Project Background\n\nBackground.\n\n'
        '## 9. Integration\n\nIntro.\n\n```\n' + ASCII_BOX + '\n```\n\n'
        '| Interface | Owner |\n|---|---|\n| POS | Client |\n| CRM | Sun* |\n'
    )
    layouts = [e['layout'] for e in parse_profile(doc).auto_extra_slides]
    assert 'diagram' in layouts, layouts
    assert 'table' in layouts, f'the chapter table left the deck: {layouts}'
