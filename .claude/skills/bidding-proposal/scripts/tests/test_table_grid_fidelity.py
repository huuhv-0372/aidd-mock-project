"""A grid on a slide must say what the grid in the markdown said.

Four defects, all on the 1.4 Feature List slides, all of them the renderer
disagreeing with the source about the shape of a table:

  * an empty cell was dropped, so `| | Total | | 36 |` printed `Total` under `#`
    and `36` under `Group` — a total of thirty-six groups, which nobody wrote;
  * columns were measured in CJK widths whatever the script, so latin text was
    judged 40% longer than it renders and cut: `Authentication & Onboa…`;
  * a flat one-inch floor gave `#` and `Count` an inch each and took it from the
    columns holding sentences;
  * every page was filled to capacity and the last took what was left, so a
    seven-row table came out as six rows and a slide holding one.
"""
from __future__ import annotations

import pytest

pytest.importorskip('pptx')

from lib.renderer import PPTXRenderer  # noqa: E402
from lib.templates.base import ContentKind, ShapeRole  # noqa: E402

SCREEN_GROUPS = '''| # | Group | Screens | Count |
|---|---|---|---|
| 1 | Authentication & Onboarding | Sign In, Sign Up, Pairing flow, Basic Info | 11 |
| 2 | Home / Dashboard | Blank home, Main home, Low-battery banner | 5 |
| 6 | Diary | Historical view | 1 |
| | Total | | 36 |
'''


def test_an_empty_cell_stays_a_cell():
    rows = PPTXRenderer._parse_md_table_rows(SCREEN_GROUPS)
    assert rows[0] == ['#', 'Group', 'Screens', 'Count']
    assert rows[-1] == ['', 'Total', '', '36'], 'Total belongs under Group, 36 under Count'
    assert all(len(row) == 4 for row in rows), 'every row keeps all four columns'


def test_a_row_of_nothing_is_dropped():
    assert PPTXRenderer._parse_md_table_rows('| | | |\n| a | b | c |') == [['a', 'b', 'c']]


def test_latin_text_is_measured_as_latin():
    """The column that held `Authentication & Onboarding` in two lines."""
    renderer = PPTXRenderer()
    latin = renderer._cell_chars_per_line(1.9, sample='Authentication & Onboarding')
    cjk = renderer._cell_chars_per_line(1.9, sample='認証とオンボーディング')
    assert latin > cjk, 'latin characters are narrower than CJK ones'
    assert latin >= len('Authentication & Onboarding') / 2, \
        'two lines of this column must hold the phrase'


def test_a_two_digit_column_is_not_given_an_inch(template_path):
    """Width follows content: the floor is the column's longest WORD."""
    from pptx import Presentation

    renderer = PPTXRenderer()
    rows = PPTXRenderer._parse_md_table_rows(SCREEN_GROUPS)
    table = next(shape.table for slide in Presentation(str(template_path)).slides
                 for shape in slide.shapes
                 if shape.has_table and len(shape.table.columns) == 4)
    total = sum(col.width for col in table.columns)
    renderer._spread_col_widths(table, rows, total)

    widths = [col.width / renderer._EMU_PER_INCH for col in table.columns]
    assert sum(col.width for col in table.columns) == total, 'the grid keeps its width'
    assert widths[0] < 0.8, f'the `#` column holds one digit: {widths}'
    assert widths[3] < 0.8, f'the `Count` column holds two digits: {widths}'
    assert widths[2] > widths[1] > widths[0], f'the sentences get the width: {widths}'
    # Every column still fits its longest word without breaking it.
    for j, width in enumerate(widths):
        longest = max((renderer._text_width_in(word, renderer._TABLE_FONT_PT)
                       for row in rows if j < len(row) for word in row[j].split()),
                      default=0.0)
        assert width >= longest, f'column {j} breaks a word: {width} < {longest}'


@pytest.mark.parametrize('rows,max_rows,expected', [
    (7, 6, 4),    # not 6 and 1
    (33, 5, 5),
    (6, 5, 3),    # not 5 and 1
    (3, 5, 5),    # one page, no balancing needed
])
def test_pages_are_filled_evenly(rows, max_rows, expected):
    assert PPTXRenderer._rows_per_page(rows, max_rows) == expected


def test_a_page_never_holds_more_than_the_slide_fits():
    for rows in range(1, 60):
        for cap in range(1, 12):
            per_page = PPTXRenderer._rows_per_page(rows, cap)
            assert per_page <= max(cap, 1)
            pages = -(-rows // per_page)
            assert pages == max(1, -(-rows // cap)), 'balancing must not add a page'


def test_a_splittable_table_is_not_trimmed_to_save_a_page(monkeypatch, template_path):
    """Full text is the point: the split pass makes the room, not the scissors."""
    from pptx import Presentation

    renderer = PPTXRenderer()
    role = ShapeRole(name='function_table', cardinality=1, kind=ContentKind.TABLE,
                     source_path='features.table')
    shape = next(shape for slide in Presentation(str(template_path)).slides
                 for shape in slide.shapes
                 if shape.has_table and len(shape.table.columns) >= 4)
    rows = [['#', 'Feature', 'Notes'],
            ['1', 'Sign in / Sign up / Sign out / Email verification',
             'Session management, token refresh']]
    trimmed = renderer._trim_table_cells_to_fit(shape, rows, role, has_header=True,
                                                rows_per_slide=1)
    assert not trimmed, 'a row that fits a page of its own keeps every word'
    assert '…' not in ' '.join(rows[1])
