"""A box that bullets its own paragraphs must mark the list itself.

Slide 4's two columns carry `buChar='●'` on every item paragraph. Writing a list
into one of them as a single paragraph of newline-separated text produced a slide
with one round bullet and five unmarked lines — and because the text also carried
its own `■` glyph, the one marked line read `● ■ item`: two markers, two shapes,
neither intended.

These tests pin the fix from both ends: the glyph is the shape's decision, and a
list becomes one paragraph per item.
"""
from __future__ import annotations

import pytest

pytest.importorskip('pptx')

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from lib.renderer import PPTXRenderer  # noqa: E402

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
SLIDE4_INDEX = 3
LEFT_BODY = 'Google Shape;762;p4'
RIGHT_BODY = 'Google Shape;765;p4'


def _shape(template_path, name: str):
    slide = Presentation(str(template_path)).slides[SLIDE4_INDEX]
    return next(s for s in slide.shapes if s.name == name)


def _bullets(text_frame) -> list:
    """Per paragraph: '●' for a native bullet, None for an explicit buNone."""
    out = []
    for para in text_frame.paragraphs:
        pPr = para._p.find(f'{A}pPr')
        if pPr is None:
            out.append('?')
            continue
        char = pPr.find(f'{A}buChar')
        out.append(char.get('char') if char is not None
                   else (None if pPr.find(f'{A}buNone') is not None else '?'))
    return out


def test_the_template_columns_really_do_bullet_themselves(template_path):
    """The premise. If this ever changes, the rest of this file is moot."""
    for name in (LEFT_BODY, RIGHT_BODY):
        assert PPTXRenderer._native_bullet(_shape(template_path, name).text_frame)


def test_a_plain_box_is_not_reported_as_bulleting_itself(template_path):
    """The cover's title has no list styling, so it must still take a glyph."""
    slide = Presentation(str(template_path)).slides[0]
    plain = [s for s in slide.shapes
             if s.has_text_frame and not PPTXRenderer._native_bullet(s.text_frame)]
    assert plain, 'at least one box on the cover bullets nothing'


def test_each_item_becomes_its_own_paragraph_with_the_templates_bullet(template_path):
    shape = _shape(template_path, LEFT_BODY)
    lines = ['First challenge', 'Second challenge', 'Third challenge']
    PPTXRenderer._set_text_lines_preserving_format(shape.text_frame, lines)
    assert [p.text for p in shape.text_frame.paragraphs] == lines
    assert _bullets(shape.text_frame) == ['●', '●', '●']


def test_no_glyph_is_written_into_the_text(template_path):
    """The `■` belonged to the drawn appendix layouts, never to this box."""
    shape = _shape(template_path, LEFT_BODY)
    PPTXRenderer._set_text_lines_preserving_format(
        shape.text_frame, ['Alpha', 'Beta', 'Gamma'])
    for para in shape.text_frame.paragraphs:
        assert '■' not in para.text
        assert '●' not in para.text, 'the marker is styling, not characters'


def test_every_line_keeps_one_run(template_path):
    """A cloned paragraph with two runs would print its line twice."""
    shape = _shape(template_path, RIGHT_BODY)
    PPTXRenderer._set_text_lines_preserving_format(shape.text_frame, ['One', 'Two'])
    for para in shape.text_frame.paragraphs:
        assert len(para.runs) == 1


def test_a_colon_lead_in_keeps_its_words_and_loses_its_marker(template_path):
    """`Concrete success indicators:` announces the items; it is not one of them."""
    shape = _shape(template_path, RIGHT_BODY)
    lines = ['Concrete success indicators:', 'Sensors report without loss.',
             'Connectivity stays stable.']
    PPTXRenderer._set_text_lines_preserving_format(shape.text_frame, lines)
    assert [p.text for p in shape.text_frame.paragraphs] == lines, 'nothing dropped'
    assert _bullets(shape.text_frame) == [None, '●', '●']


def test_a_trailing_colon_on_the_LAST_line_stays_bulleted(template_path):
    """It introduces nothing, so it is an item like any other."""
    shape = _shape(template_path, RIGHT_BODY)
    PPTXRenderer._set_text_lines_preserving_format(
        shape.text_frame, ['A real item.', 'Ends with a colon:'])
    assert _bullets(shape.text_frame) == ['●', '●']


def test_a_single_line_is_written_as_one_paragraph(template_path):
    shape = _shape(template_path, LEFT_BODY)
    PPTXRenderer._set_text_lines_preserving_format(shape.text_frame, ['Only one'])
    assert len(shape.text_frame.paragraphs) == 1
    assert shape.text_frame.paragraphs[0].text == 'Only one'


def test_a_box_with_no_runs_falls_back_without_raising():
    """An empty text frame has no model paragraph to clone."""
    prs = Presentation()
    box = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(1), Inches(3), Inches(1))
    PPTXRenderer._set_text_lines_preserving_format(box.text_frame, ['A', 'B'])
    assert 'A' in box.text_frame.text and 'B' in box.text_frame.text
