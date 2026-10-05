"""Container chapter (schedule + deliverables) and divider-only `team` section:
one divider named after the source chapter, the Gantt drawn natively into
slide 43, 体制 extras pinned after divider 17, template caption stripped."""
from __future__ import annotations

import pytest

from lib.profile_parser import parse_profile
from lib.renderer import PPTXRenderer
from tests.test_renderer_slide_selection import slide_texts, _walk

pptx = pytest.importorskip('pptx')
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

ORG = {'layout': 'org_chart', 'title': 'プロジェクト体制（要件定義）', 'anchor_section': 'team',
       'lanes': [{'title': '貴社', 'accent': 'gold', 'nodes': [{'label': '業務部長'}]},
                 {'title': 'Sun*', 'accent': 'red', 'nodes': [{'label': 'PSM'}]}],
       'roles': [{'name': 'PSM', 'desc': '全体管理'}]}
ORG2 = {**ORG, 'title': 'プロジェクト体制（開発・テスト）'}


@pytest.fixture(scope='module')
def profile(fixtures_dir):
    return parse_profile((fixtures_dir / 'jp-schedule-gantt-team-container.md')
                         .read_text(encoding='utf-8'))


@pytest.fixture(scope='module')
def deck(profile, template_path, tmp_path_factory):
    out = PPTXRenderer().render_from_profile(
        profile, str(template_path), 'out', str(tmp_path_factory.mktemp('deck')),
        extra_slides=[ORG, ORG2] + list(profile.auto_extra_slides))
    return pptx.Presentation(out)


def test_parser_reads_gantt_rows_milestones_and_container(profile):
    assert [r['owner'] for r in profile.schedule.gantt_rows] == ['sun', 'client', 'sun', 'client', 'next']
    assert profile.schedule.gantt_rows[3]['end'] == '2027/08'
    assert [m['label'] for m in profile.schedule.milestones] == ['★P1稼働', '▲要件定義完了']
    assert profile.schedule.milestones[1]['at'] == '2027/03'
    assert profile.schedule.table == []  # the Gantt table is not also an orphan table
    assert profile.section_chapters['schedule'] == profile.section_chapters['deliverables'] \
        == '開発スケジュールおよび成果物'
    assert profile.section_order.index('schedule') < profile.section_order.index('deliverables')
    assert profile.team.description.startswith('要件定義フェーズ') and profile.team.rows == []
    assert not profile.auto_extra_slides


def test_one_divider_named_after_the_container_and_schedule_before_deliverables(deck):
    texts = slide_texts(deck)
    dividers = [i for i, t in enumerate(texts)
                if t.split() == ['Schedule', '&', 'Deliverables', '開発スケジュールおよび成果物']]
    assert len(dividers) == 1
    assert not any('想定作成物' in t or '想定成果物' in t or '実施スケジュール' in t for t in texts)
    schedule = next(i for i, t in enumerate(texts) if '全体スケジュール案' in t)
    deliverables = next(i for i, t in enumerate(texts) if '提出物' in t and '要件定義書' in t)
    assert dividers[0] < schedule < deliverables
    assert 'Schedule & Deliverables' in texts[schedule] and 'Schedule & Deliverables' in texts[deliverables]


def test_gantt_is_drawn_natively_on_the_schedule_slide(deck):
    texts = slide_texts(deck)
    slide = deck.slides[next(i for i, t in enumerate(texts) if '全体スケジュール案' in t)]
    # The template's legend pills are rounded rectangles too — bars are the
    # text-less ones inside the frame.
    bars = [s for s in slide.shapes if s.shape_type == 1 and s.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
            and not s.text_frame.text.strip() and s.height < 914400 * 0.3]
    markers = [s for s in slide.shapes if s.shape_type == 1 and s.auto_shape_type == MSO_SHAPE.ISOSCELES_TRIANGLE]
    assert len(bars) == 5 and len(markers) == 2
    assert not any(getattr(s, 'has_table', False) and s.has_table for s in slide.shapes)
    joined = texts[deck.slides.index(slide)]
    assert '★P1稼働' in joined and 'P1 要件定義' in joined
    assert 'スプリント開発' not in joined and 'キャッチアップ' not in joined  # template sample boxes gone


def test_team_divider_is_kept_renamed_and_followed_by_its_extras(deck):
    texts = slide_texts(deck)
    divider = next(i for i, t in enumerate(texts)
                   if t.split() == ['Project', 'Structure', 'プロジェクト体制'])
    assert not any('Sun*のスコープ' in t or 'Sun*’s Scope' in t for t in texts)
    assert 'プロジェクト体制（要件定義）' in texts[divider + 1]
    assert 'プロジェクト体制（開発・テスト）' in texts[divider + 2]
    assert 'Project Structure' in texts[divider + 1]  # breadcrumb from the AGENDA gloss
    cost = next(i for i, t in enumerate(texts) if 'PSM' in t and '7.5' in t)
    assert divider < cost


def test_template_sample_caption_is_stripped(deck):
    assert not any('機能一覧表サンプル' in t for t in slide_texts(deck))


def test_team_divider_survives_without_a_markdown_chapter(fixtures_dir, template_path, tmp_path):
    """The body of `team` may be ONLY anchored extras (SKILL.md: divider, no
    content slide). With the `## 6. プロジェクト体制` chapter removed, the divider
    must still be kept, named by its default, and followed by the org charts."""
    md = (fixtures_dir / 'jp-schedule-gantt-team-container.md').read_text(encoding='utf-8')
    head, _, rest = md.partition('## 6. プロジェクト体制')
    md = head + '## 費用' + rest.partition('## 費用')[2]
    md = md.replace('  "6. プロジェクト体制": team\n', '').replace('- プロジェクト体制 — Project Structure\n', '')
    profile = parse_profile(md)
    assert 'team' not in profile.section_order
    out = PPTXRenderer().render_from_profile(profile, str(template_path), 'out2', str(tmp_path),
                                             extra_slides=[ORG, ORG2])
    texts = slide_texts(pptx.Presentation(out))
    divider = next(i for i, t in enumerate(texts) if t.split() == ['Project', 'Structure', 'プロジェクト体制'])
    assert 'プロジェクト体制（要件定義）' in texts[divider + 1]
    assert 'プロジェクト体制（開発・テスト）' in texts[divider + 2]
    assert not any('Appendix' in t for t in texts[divider + 1:divider + 3])


def test_schema_headings_do_not_retitle_a_japanese_deck():
    """`## Features` is gen-md plumbing: on a Japanese document the template's
    機能一覧 title must survive; a heading in the author's own words still wins."""
    md = ('# Project Profile: x\n\n## Features\n\n### Description\n機能です。\n\n'
          '| No | 機能 |\n|---|---|\n| 1 | せり |\n\n## 全体スケジュール案\n2026年12月着手。\n')
    profile = parse_profile(md)
    assert profile.is_japanese_source
    assert 'features' not in profile.section_headings
    assert profile.section_headings['schedule'] == '全体スケジュール案'
    english = parse_profile('# Project Profile: y\n\n## Features\n\n### Description\nAuction.\n')
    assert english.section_headings['features'] == 'Features'  # an English deck keeps English
