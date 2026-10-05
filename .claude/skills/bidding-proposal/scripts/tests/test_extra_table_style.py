"""Pin the appendix table's styling to the template's own tables.

`extra_slide_layouts.draw_table_in_box` claims its colors, borders and margins
were copied verbatim off the template's 機能一覧表 (slide 5). This turns that
claim into a check: the constants are read back out of the .pptx, so editing
the template — or drifting the constants toward the generate-slide palette the
appendix tables used to wear (#B69256 / #DDDDDD) — fails here instead of
shipping a deck whose appendix tables sit a shade off every other table in it.

Slide 5 is the reference because it is the table `SECTION_TO_SLIDES['features']`
fills, so a reader meets it before any appendix slide.
"""
from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.util import Emu

from lib.extra_slide_layouts import (
    TBL_BODY_FILL,
    TBL_BORDER,
    TBL_BORDER_PT,
    TBL_HEAD_FILL,
    TBL_MARGIN_X_IN,
)

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
REFERENCE_SLIDE = 5  # 機能一覧表 — see module docstring


@pytest.fixture(scope='module')
def reference_table(template_path):
    prs = Presentation(str(template_path))
    for shape in prs.slides[REFERENCE_SLIDE - 1].shapes:
        if shape.has_table:
            return shape.table
    pytest.fail(f'template slide {REFERENCE_SLIDE} has no table')


def _cell_fill(cell) -> str:
    srgb = cell._tc.find(f'{A}tcPr/{A}solidFill/{A}srgbClr')
    assert srgb is not None, 'reference cell has no explicit fill'
    return srgb.get('val').upper()


def _cell_border(cell, side: str) -> tuple[int, str]:
    line = cell._tc.find(f'{A}tcPr/{A}{side}')
    assert line is not None, f'reference cell has no {side}'
    srgb = line.find(f'{A}solidFill/{A}srgbClr')
    return int(line.get('w')), srgb.get('val').upper()


def test_header_fill_matches_template(reference_table):
    assert _cell_fill(reference_table.cell(0, 0)) == str(TBL_HEAD_FILL)


def test_body_fill_matches_template(reference_table):
    assert _cell_fill(reference_table.cell(1, 0)) == str(TBL_BODY_FILL)


@pytest.mark.parametrize('side', ['lnL', 'lnR', 'lnT', 'lnB'])
def test_borders_match_template(reference_table, side):
    width_emu, color = _cell_border(reference_table.cell(1, 0), side)
    assert color == str(TBL_BORDER)
    # 12700 EMU to the point; a hair of rounding either way is the template's,
    # not a different weight.
    assert abs(width_emu - TBL_BORDER_PT * 12700) < 30


def test_horizontal_margin_matches_template(reference_table):
    tcPr = reference_table.cell(1, 0)._tc.find(f'{A}tcPr')
    assert abs(int(tcPr.get('marL')) - TBL_MARGIN_X_IN * 914400) < 400


def test_cells_are_vertically_centred(reference_table):
    tcPr = reference_table.cell(1, 0)._tc.find(f'{A}tcPr')
    assert tcPr.get('anchor') == 'ctr'


def test_drawn_table_carries_no_theme_style(template_path, tmp_path):
    """A drawn table must ship the bare `a:tblPr` every template table has.

    python-pptx stamps `firstRow`/`bandRow` and a `tableStyleId` on every new
    table, which paints a header band and zebra striping UNDER the explicit
    per-cell fills — banding no table in this template has.
    """
    from lib.extra_slide_layouts import draw_table_in_box

    prs = Presentation(str(template_path))
    slide = prs.slides[REFERENCE_SLIDE - 1]
    before = {id(sh) for sh in slide.shapes}
    draw_table_in_box(
        slide, Emu(914400), Emu(914400), Emu(4572000), Emu(1828800),
        ['A', 'B'], [['1', '2'], ['3', '4']],
    )
    drawn = next(sh for sh in slide.shapes if id(sh) not in before and sh.has_table)

    tblPr = drawn.table._tbl.find(f'{A}tblPr')
    assert tblPr is None or (
        not tblPr.findall(f'{A}tableStyleId')
        and not {'firstRow', 'bandRow'} & set(tblPr.attrib)
    )
    assert _cell_fill(drawn.table.cell(0, 0)) == str(TBL_HEAD_FILL)
    assert _cell_border(drawn.table.cell(1, 0), 'lnT')[1] == str(TBL_BORDER)
