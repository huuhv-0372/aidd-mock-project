"""The template-contract gate, run as a test.

`lib/templates/svn.py` addresses the shipped .pptx by exact string and by shape
index. Editing the template silently disables a mapping instead of erroring, so
this is what turns that into a failing build.
"""
from __future__ import annotations

import os
import subprocess
import sys

import pytest

pytest.importorskip('pptx')


def test_template_contract_holds(scripts_dir, template_path):
    result = subprocess.run(
        [sys.executable, str(scripts_dir / 'verify-template-contract.py')],
        capture_output=True, text=True, cwd=str(scripts_dir),
    )
    assert result.returncode == 0, (
        'A svn.py mapping no longer resolves against the shipped template.\n'
        'Fix the mapping or the template — do NOT delete the entry to make this '
        'pass; dropping a TEMPLATE_NOTE_TEXTS line re-ships a Vietnamese '
        f'authoring note to the client.\n\n{result.stdout}\n{result.stderr}'
    )


def test_verifier_detects_a_broken_mapping(scripts_dir, template_path, tmp_path):
    """The gate must actually fail when a mapping goes stale.

    Guards against the check quietly degrading into a no-op — a green gate that
    cannot go red is worse than no gate.
    """
    script = (scripts_dir / 'verify-template-contract.py').read_text(encoding='utf-8')
    probe = tmp_path / 'probe.py'
    probe.write_text(
        script.replace(
            'check_slide_text_map(report, prs, \'SLIDE_TEXT_OVERRIDES\', SLIDE_TEXT_OVERRIDES)',
            "check_slide_text_map(report, prs, 'SLIDE_TEXT_OVERRIDES', "
            "{21: {'text that is not in the template': 'x'}})",
        ),
        encoding='utf-8',
    )
    # The probe lives in tmp_path, so its own `sys.path.insert(parent)` points
    # away from the skill — hand it the scripts dir explicitly.
    # …and its template resolution is relative to that same parent, so name the
    # shipped template explicitly.
    env = {**os.environ, 'PYTHONPATH': str(scripts_dir)}
    result = subprocess.run(
        [sys.executable, str(probe), '--template', str(template_path)],
        capture_output=True, text=True, cwd=str(scripts_dir), env=env)
    assert result.returncode == 1, f'{result.stdout}\n{result.stderr}'
    assert 'unmatched' in result.stdout


def test_every_role_source_names_a_section_the_deck_knows():
    """An overflow slide must be able to find its way back to its chapter.

    `PPTXRenderer._queue_overflow_extra` reads the section out of the role's
    `source_path`, and `_resolve_extra_anchor` places the slide after that
    section. A source naming nothing the deck knows resolves to no anchor, and
    the slide is appended to the END of the deck — on a real proposal that put
    the continuation of slide 23's cards on slide 55, thirty-two slides from the
    chapter it belongs to.

    `SOURCE_SECTION_ALIASES` is the translation for the roles whose source is
    not named after its own section. This test is what makes adding one
    non-optional when a new role is wired up.
    """
    from lib.templates.svn import (
        SECTION_TO_SLIDES,
        SOURCE_SECTION_ALIASES,
        SVNProposalTemplate,
    )

    unresolvable = []
    for slide_no, config in SVNProposalTemplate().slide_role_configs.items():
        for role in config.roles:
            root = (role.source_path or '').split('.')[0].split('[')[0]
            if not root:
                continue
            if SOURCE_SECTION_ALIASES.get(root, root) not in SECTION_TO_SLIDES:
                unresolvable.append(f'slide {slide_no} role {role.name!r} -> {root!r}')

    assert unresolvable == [], (
        'These role sources name no section in SECTION_TO_SLIDES, so any text '
        'that overflows their box lands at the end of the deck instead of after '
        'its own chapter. Add an entry to SOURCE_SECTION_ALIASES in svn.py:\n  '
        + '\n  '.join(unresolvable)
    )


def test_no_template_sample_copy_survives_a_render(tmp_path, template_path,
                                                   fixtures_dir):
    """No deck may ship the sample sentences the template was cloned with.

    Slides 33, 34, 36, 43 and 56 carry copy written for another proposal in
    shapes no role is bound to: a lead-in about somebody else's architecture, a
    claim that the client supplied an NFR document, and the whole of another
    project's Gantt chart. They rendered into a deck written entirely in English
    as Japanese paragraphs under English headings.
    """
    import io
    from contextlib import redirect_stdout

    from pptx import Presentation

    from lib.profile_parser import parse_profile
    from lib.renderer import PPTXRenderer
    from lib.templates.svn import TEMPLATE_SAMPLE_TEXTS

    def walk(shapes):
        for shape in shapes:
            yield shape
            if 'GROUP' in str(shape.shape_type):
                yield from walk(shape.shapes)

    leaked = []
    for fixture in sorted(fixtures_dir.glob('*.md')):
        profile = parse_profile(fixture.read_text(encoding='utf-8'))
        with redirect_stdout(io.StringIO()):
            out = PPTXRenderer().render_from_profile(
                profile, str(template_path), fixture.stem, str(tmp_path),
                extra_slides=list(profile.auto_extra_slides))
        prs = Presentation(str(out))
        for number, slide in enumerate(prs.slides, 1):
            for shape in walk(slide.shapes):
                if not shape.has_text_frame:
                    continue
                text = '\n'.join(
                    line for line in (raw.strip() for raw
                                      in shape.text_frame.text.splitlines()) if line)
                if text in TEMPLATE_SAMPLE_TEXTS:
                    leaked.append(f'{fixture.stem} slide {number}: {text[:50]!r}')

    assert leaked == [], 'template sample copy reached the deck:\n  ' + '\n  '.join(leaked)
