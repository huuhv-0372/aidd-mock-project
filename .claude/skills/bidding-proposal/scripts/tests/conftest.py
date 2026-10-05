"""Shared pytest fixtures for the bidding-proposal suite.

Puts `scripts/` on `sys.path` the same way `gen-slide.py` does, so tests import
`lib.*` exactly as the entry points do rather than through a package install.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

SCRIPTS_DIR = Path(__file__).resolve().parent.parent
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))


@pytest.fixture(scope='session')
def scripts_dir() -> Path:
    return SCRIPTS_DIR


@pytest.fixture(scope='session')
def fixtures_dir() -> Path:
    return Path(__file__).resolve().parent / 'fixtures'


@pytest.fixture(scope='session')
def template_path(scripts_dir: Path) -> Path:
    """The SVN template that ships with the skill."""
    return scripts_dir.parent / 'templates' / 'SVN Proposal Menu.pptx'


def read_fixture(fixtures_dir: Path, name: str) -> str:
    return (fixtures_dir / name).read_text(encoding='utf-8')


@pytest.fixture(scope='session')
def golden_dir() -> Path:
    return Path(__file__).resolve().parent / 'fixtures' / 'golden'


def render_and_extract(fixture: Path, template: Path, out_dir) -> dict:
    """Render *fixture* and return {slide number: [text, ...]}.

    Text, not bytes: a .pptx carries timestamps and part ordering that differ
    between two runs of identical code, so a byte comparison would fail for
    reasons no reader could see. What a golden file must protect is what the
    deck SAYS and how many slides it takes to say it.

    Rendered exactly the way `gen-slide.py` renders — auto-routed chapters
    included — so the golden covers the path users actually take.
    """
    import sys as _sys
    if str(SCRIPTS_DIR) not in _sys.path:
        _sys.path.insert(0, str(SCRIPTS_DIR))
    from pptx import Presentation

    from lib.profile_parser import parse_profile
    from lib.renderer import PPTXRenderer

    profile = parse_profile(fixture.read_text(encoding='utf-8'))
    out = PPTXRenderer().render_from_profile(
        profile, str(template), fixture.stem, str(out_dir),
        extra_slides=list(profile.auto_extra_slides),
    )
    prs = Presentation(str(out))

    def walk(shapes):
        for shape in shapes:
            yield shape
            if 'GROUP' in str(shape.shape_type):
                yield from walk(shape.shapes)

    return {
        str(i): [' '.join(s.text_frame.text.split())
                 for s in walk(slide.shapes)
                 if s.has_text_frame and s.text_frame.text.strip()]
        for i, slide in enumerate(prs.slides, 1)
    }
