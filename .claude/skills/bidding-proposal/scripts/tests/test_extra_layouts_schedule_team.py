"""`gantt` and `org_chart` extra-slide layouts render the shapes they promise,
inside the slide, and refuse (loudly) what they cannot fit."""
from __future__ import annotations

import json

import pytest

from lib.extra_slide_layouts import _LAYOUTS, render_extra_layout
from lib.gantt_layout import gantt_rows, owner_key

pptx = pytest.importorskip('pptx')
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402


@pytest.fixture
def entries(fixtures_dir):
    return json.loads((fixtures_dir / 'extra-gantt-org-chart.json').read_text(encoding='utf-8'))


@pytest.fixture
def blank_slide(template_path):
    prs = pptx.Presentation(str(template_path))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return prs, slide


def _autoshapes(slide, kind):
    return [s for s in slide.shapes
            if s.shape_type == 1 and s.auto_shape_type == kind]  # 1 = AUTO_SHAPE


def _inside(prs, slide):
    return all(0 <= s.left and 0 <= s.top and s.left + s.width <= prs.slide_width + 1
               and s.top + s.height <= prs.slide_height + 1 for s in slide.shapes)


def test_layouts_registered():
    assert {'gantt', 'org_chart'} <= set(_LAYOUTS)


def test_gantt_draws_one_bar_per_row_and_one_marker_per_milestone(entries, blank_slide):
    prs, slide = blank_slide
    entry = entries[0]
    render_extra_layout(slide, entry, prs.slide_width, prs.slide_height)
    bars = _autoshapes(slide, MSO_SHAPE.ROUNDED_RECTANGLE)
    assert len(bars) == len(entry['rows'])
    assert len(_autoshapes(slide, MSO_SHAPE.ISOSCELES_TRIANGLE)) == len(entry['milestones'])
    texts = ' '.join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert '26/12' in texts and '27/11' in texts and '★P1稼働' in texts
    assert '貴社' in texts and '次フェーズ' in texts  # legend
    assert _inside(prs, slide)


def test_gantt_owner_mapping_and_axis_derivation(entries):
    rows, milestones, months = gantt_rows(entries[0])
    assert [r['owner'] for r in rows][:4] == ['sun', 'client', 'client', 'sun']
    assert rows[-1]['owner'] == 'next'
    assert months[0] == '2026/12' and months[-1] == '2027/11' and len(months) == 12
    assert owner_key('貴社') == 'client' and owner_key('P2') == 'next' and owner_key('') == 'sun'
    assert owner_key('外部ベンダー') == 'other'


@pytest.mark.parametrize('bad', [
    {'rows': []},
    {'rows': [{'task': 't', 'start': '2026-1x'}]},
    {'rows': [{'task': f't{i}', 'start': '2026/01', 'end': '2026/02'} for i in range(11)]},
    {'rows': [{'task': 't', 'start': '2026/01', 'end': '2027/06'}]},  # 18 months
    {'rows': [{'task': 't', 'start': '2026/05', 'end': '2026/02'}]},  # ends before it starts
])
def test_gantt_refuses_bad_or_oversized_input(bad, blank_slide):
    prs, slide = blank_slide
    with pytest.raises(ValueError):
        render_extra_layout(slide, {'layout': 'gantt', 'title': 'x', **bad},
                            prs.slide_width, prs.slide_height)


def test_org_chart_draws_every_node_and_role(entries, blank_slide):
    prs, slide = blank_slide
    entry = entries[1]
    render_extra_layout(slide, entry, prs.slide_width, prs.slide_height)
    texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    for lane in entry['lanes']:
        assert lane['title'] in texts
        for node in lane['nodes']:
            assert node['label'] in texts
    roles_box = next(t for t in texts if t.startswith('・PSM'))
    assert roles_box.count('\n') + 1 == len(entry['roles'])
    assert '各ロールの役割' in texts
    assert _inside(prs, slide)


def test_org_chart_caps(blank_slide):
    prs, slide = blank_slide
    too_many_nodes = {'layout': 'org_chart', 'title': 'x',
                      'lanes': [{'title': 'a', 'nodes': [{'label': str(i)} for i in range(7)]}]}
    with pytest.raises(ValueError):
        render_extra_layout(slide, too_many_nodes, prs.slide_width, prs.slide_height)
    four_lanes = {'layout': 'org_chart', 'title': 'x', 'lanes': [{'title': str(i), 'nodes': []} for i in range(4)]}
    with pytest.raises(ValueError):
        render_extra_layout(slide, four_lanes, prs.slide_width, prs.slide_height)
