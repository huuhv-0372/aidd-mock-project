"""A chapter holding several labelled tables gets a slide, and a label, per table.

`### 1.4 Feature List` nests its real material in `#### ` sub-blocks — screen
groups, the full feature breakdown, the out-of-scope list — each with its own
columns. Read as one block they merged: the second table's header row rendered as
a data row of the first, its surplus column was cut away to fit the first table's
grid, and six continuation slides all wore the one label.

Slide 5 has a pill above its grid whose only job is to name the table below it.
These tests pin that it does.
"""
from __future__ import annotations

import pytest

pytest.importorskip('pptx')

from pptx import Presentation  # noqa: E402

from lib.markdown_primitives import parse_prose, split_sub_subsections  # noqa: E402
from lib.profile_parser import parse_profile  # noqa: E402
from lib.renderer import PPTXRenderer  # noqa: E402
from lib.templates.svn import SVNProposalTemplate  # noqa: E402

PILL = 'Google Shape;776;p5'

DOC = '''---
sections:
  "1.4. Feature List": features
---

# Proposal

## Agenda

1. Features

## 1.4. Feature List

The MVP covers 12 screens in 3 areas. The tables below list every unit.

#### 1.4.1. Screen groups

| # | Group | Screens | Count |
|---|---|---|---|
| 1 | Authentication | Sign in, Sign up | 2 |
| 2 | Home | Dashboard | 1 |

#### 1.4.2. Full feature breakdown

| # | Category | Feature | Notes | Priority |
|---|---|---|---|---|
| 1 | Auth | Email sign-in | Session handling | Must |
| 2 | Home | Metric cards | Six cards | Must |

#### 1.4.3. Out-of-scope features (confirmed)

- Admin console.
- Theme switching.
'''

SINGLE_TABLE_DOC = '''# Proposal

## Features

Every unit the team estimated.

| # | Feature |
|---|---|
| 1 | Sign in |
| 2 | Dashboard |
'''


# --- parsing ---------------------------------------------------------------

def test_each_sub_block_with_a_table_becomes_its_own_labelled_group():
    features = parse_profile(DOC).features
    assert [g['label'] for g in features.groups] == [
        'Screen groups', 'Full feature breakdown']


def test_the_second_tables_columns_survive_intact():
    """Merged into the first table's grid, the Priority column was cut away."""
    groups = parse_profile(DOC).features.groups
    assert list(groups[0]['table'][0]) == ['#', 'Group', 'Screens', 'Count']
    assert list(groups[1]['table'][0]) == ['#', 'Category', 'Feature', 'Notes', 'Priority']


def test_no_tables_header_row_becomes_another_tables_data_row():
    for group in parse_profile(DOC).features.groups:
        for row in group['table']:
            assert 'Category' not in row.values(), 'a header leaked in as data'


def test_the_description_carries_no_heading_lines():
    """The complaint this started from: `#### 1.4.1. Screen groups` as loose words."""
    description = parse_profile(DOC).features.description
    assert '####' not in description
    assert 'Screen groups' not in description
    assert 'Full feature breakdown' not in description
    assert 'The MVP covers 12 screens' in description


def test_parse_prose_drops_headings_at_every_level():
    assert parse_prose('# One\n\nBody.\n\n###### Six\n\nMore.') == 'Body.\n\nMore.'
    assert '#hashtag' in parse_prose('#hashtag is not a heading'), 'needs a space'


def test_a_tableless_sub_block_reaches_the_deck_on_its_own_slide():
    """The out-of-scope list fits no grid, and must not vanish because of it."""
    profile = parse_profile(DOC)
    titles = [e.get('title', '') for e in profile.auto_extra_slides]
    assert any('Out-of-scope' in t for t in titles)
    entry = next(e for e in profile.auto_extra_slides if 'Out-of-scope' in e['title'])
    assert 'Admin console.' in entry.get('bullets', [])


def test_a_single_unlabelled_table_still_fills_the_slide():
    """The shape gen-md.py emits: one group, no label, behaviour unchanged."""
    features = parse_profile(SINGLE_TABLE_DOC).features
    assert len(features.groups) == 1
    assert features.groups[0]['label'] == '', 'no label means the pill keeps its own words'
    assert features.groups[0]['table'] == features.table


def test_split_sub_subsections_reads_only_the_fourth_level():
    parts = split_sub_subsections('## Two\n\n### Three\n\n#### Four\n\nbody\n')
    assert list(parts) == ['Four']
    assert parts['Four'] == 'body'


# --- rendering -------------------------------------------------------------

def test_the_role_knows_which_shape_labels_its_table():
    role = next(r for r in SVNProposalTemplate().slide_role_configs[5].roles
                if r.name == 'function_table')
    assert role.source_path == 'features.groups'
    assert role.label_shape_name == PILL


def test_grouped_and_flat_table_data_are_told_apart():
    assert PPTXRenderer._is_labelled_table_set([{'label': 'A', 'table': [{'x': '1'}]}])
    assert not PPTXRenderer._is_labelled_table_set([{'#': '1', 'Group': 'Auth'}])
    assert not PPTXRenderer._is_labelled_table_set([])


def test_every_slide_of_the_chapter_is_labelled_with_its_own_table(tmp_path,
                                                                   template_path):
    profile = parse_profile(DOC)
    out = PPTXRenderer().render_from_profile(
        profile, str(template_path), 'labelled', str(tmp_path),
        extra_slides=list(profile.auto_extra_slides))
    prs = Presentation(str(out))

    seen = []
    for slide in prs.slides:
        pill = next((s for s in slide.shapes if s.name == PILL), None)
        table = next((s for s in slide.shapes
                      if getattr(s, 'has_table', False) and s.has_table), None)
        if pill is not None and table is not None:
            seen.append((pill.text_frame.text.strip(),
                         [c.text for c in table.table.rows[0].cells]))

    assert [label for label, _ in seen] == ['Screen groups', 'Full feature breakdown']
    assert seen[0][1] == ['#', 'Group', 'Screens', 'Count']
    assert seen[1][1] == ['#', 'Category', 'Feature', 'Notes', 'Priority']


# --- A split table names its pages ------------------------------------------

def test_each_page_of_a_split_table_says_which_page_it_is(tmp_path, template_path):
    """One label repeated seven times told a reader nothing; `(2/7)` does."""
    from pptx import Presentation

    from lib.profile_parser import parse_profile
    from lib.renderer import PPTXRenderer

    rows = '\n'.join(
        f'| {i} | Category {i} | A feature described at a realistic length, '
        f'long enough to wrap | A note of similar length about it | Must |'
        for i in range(1, 19))
    doc = ('# P\n\n## Agenda\n\n1. Feature List\n\n## 1. Feature List\n\n'
           'The list below covers every unit.\n\n#### 1.1. Full feature breakdown\n\n'
           '| # | Category | Feature | Notes | Priority |\n|---|---|---|---|---|\n'
           + rows + '\n')
    profile = parse_profile(doc)
    out = PPTXRenderer().render_from_profile(
        profile, str(template_path), 'pages', str(tmp_path),
        extra_slides=list(profile.auto_extra_slides))

    pills = [' '.join(shape.text_frame.text.split())
             for slide in Presentation(str(out)).slides
             for shape in slide.shapes
             if shape.name == 'Google Shape;776;p5' and shape.has_text_frame]
    numbered = [pill for pill in pills if 'Full feature breakdown' in pill]
    assert len(numbered) > 1, 'this table must span several slides'
    total = len(numbered)
    assert numbered == [f'Full feature breakdown ({i}/{total})'
                        for i in range(1, total + 1)]


def test_a_continuation_page_does_not_repeat_the_lead_in(tmp_path, template_path):
    from pptx import Presentation

    from lib.profile_parser import parse_profile
    from lib.renderer import PPTXRenderer

    rows = '\n'.join(
        f'| {i} | Category {i} | A feature described at a realistic length, '
        f'long enough to wrap | A note of similar length about it | Must |'
        for i in range(1, 19))
    lead_in = 'The list below covers every unit.'
    doc = ('# P\n\n## Agenda\n\n1. Feature List\n\n## 1. Feature List\n\n'
           f'{lead_in}\n\n#### 1.1. Full feature breakdown\n\n'
           '| # | Category | Feature | Notes | Priority |\n|---|---|---|---|---|\n'
           + rows + '\n')
    profile = parse_profile(doc)
    out = PPTXRenderer().render_from_profile(
        profile, str(template_path), 'lead-in', str(tmp_path),
        extra_slides=list(profile.auto_extra_slides))

    said = sum(
        1 for slide in Presentation(str(out)).slides for shape in slide.shapes
        if shape.has_text_frame and lead_in in ' '.join(shape.text_frame.text.split()))
    assert said == 1, f'the lead-in belongs on the first page only, said {said} times'
