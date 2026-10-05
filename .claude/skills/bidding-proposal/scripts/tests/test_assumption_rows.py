"""前提条件 as a grid: one row per subsection, both columns the document's own.

Three defects met on this one slide, and each is pinned here:

  * the chapter did not resolve at all, so its nine subsections became eight
    appendix slides and one hijacked section (see `test_chapter_vocabulary.py`);
  * the row's content went in as raw markdown, and the renderer escaped the
    pipes: `&#124; Item &#124; Assumption`;
  * the label column was left to the template, printing 開発言語 / 外部API beside
    an English proposal's own category names.
"""
from __future__ import annotations

import pytest

pytest.importorskip('pptx')

from lib.profile_parser import parse_profile  # noqa: E402
from lib.templates.svn import SVNProposalTemplate  # noqa: E402

DOC = '''# Proposal

## Agenda

1. Estimation Assumptions

## 5. Estimation Assumptions

### 5.1. Development approach

| Item | Assumption |
|---|---|
| Methodology | Scrum, 2-week sprints |
| Repository | GitHub |

### 5.2. Programming stack (proposed — to be agreed with the client)

| Layer | Proposed |
|---|---|
| Mobile | React Native |
| Backend | NestJS |

If the client prefers Flutter, we will re-estimate.

### 5.3. Design

- The client's wireframes are used as-is.
- No dedicated designer role is budgeted.
'''


@pytest.fixture(scope='module')
def rows():
    return parse_profile(DOC).assumptions


def test_the_chapter_fills_its_own_slides_rather_than_being_taken_apart(rows):
    assert [row.label for row in rows] == [
        'Development approach', 'Programming stack', 'Design']


def test_a_row_never_carries_markdown_into_its_cell(rows):
    for row in rows:
        assert '|' not in row.content, row.content
        assert '---' not in row.content
        assert not row.content.lstrip().startswith('- ')


def test_a_matrix_row_reads_as_the_prose_it_always_was(rows):
    assert 'Methodology — Scrum, 2-week sprints' in rows[0].content
    assert 'Repository — GitHub' in rows[0].content


def test_a_row_keeps_its_matter_AND_the_caveat_after_it(rows):
    """`_field_text` must choose one; a grid row can hold both — table first."""
    stack = rows[1].content
    assert stack.index('Mobile — React Native') < stack.index('If the client prefers')
    assert 'Backend — NestJS' in stack


def test_a_bulleted_row_becomes_sentences(rows):
    assert rows[2].content.startswith("The client's wireframes are used as-is.")
    assert 'No dedicated designer role is budgeted.' in rows[2].content


def test_both_columns_of_the_grid_are_written():
    """Only the body column was, which left the template naming the rows."""
    for slide_no in (23, 24, 25):
        roles = SVNProposalTemplate().slide_role_configs[slide_no].roles
        assert [role.fill_cols for role in roles] == [[0, 1]], slide_no


def test_the_rendered_grid_says_what_the_document_said(tmp_path, template_path):
    from pptx import Presentation

    from lib.renderer import PPTXRenderer

    profile = parse_profile(DOC)
    out = PPTXRenderer().render_from_profile(
        profile, str(template_path), 'assumptions', str(tmp_path),
        extra_slides=list(profile.auto_extra_slides))

    grids = [shape.table for slide in Presentation(str(out)).slides
             for shape in slide.shapes if shape.has_table]
    assert grids, 'the assumptions grid must reach the deck'
    cells = [[' '.join(cell.text.split()) for cell in row.cells]
             for grid in grids for row in grid.rows]
    labels = [row[0] for row in cells]
    assert 'Development approach' in labels
    assert not any('開発方針' in label or '開発言語' in label for label in labels), \
        'a row must not be named after a category this document never used'
    for row in cells:
        assert '&#124;' not in row[1]
        # Nothing is cut short while the row still has height to spare.
        assert not row[1].endswith('…'), row[1]
